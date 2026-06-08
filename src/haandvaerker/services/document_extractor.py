"""Extract text from PDF, DOCX, XLSX, TXT, and images (OCR).

All dependencies are optional — returns empty string + logs warning when a
library is missing, so the import pipeline degrades gracefully.
"""
from __future__ import annotations
import logging
import pathlib
from typing import Optional

logger = logging.getLogger(__name__)

_SUPPORTED = {".pdf", ".docx", ".doc", ".xlsx", ".xls", ".txt", ".png", ".jpg", ".jpeg", ".tiff", ".bmp"}


def extract_text(file_path: str) -> str:
    """Return plain text from *file_path*. Raises FileNotFoundError if missing."""
    path = pathlib.Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix in {".docx", ".doc"}:
        return _extract_docx(path)
    if suffix in {".xlsx", ".xls"}:
        return _extract_xlsx(path)
    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix in {".png", ".jpg", ".jpeg", ".tiff", ".bmp"}:
        return _extract_image_ocr(path)
    logger.warning("Unsupported file type: %s", suffix)
    return ""


def supported_suffixes() -> frozenset[str]:
    return frozenset(_SUPPORTED)


# ── format-specific helpers ──────────────────────────────────────────────────

def _extract_pdf(path: pathlib.Path) -> str:
    try:
        import pdfplumber  # type: ignore
    except ImportError:
        logger.warning("pdfplumber not installed — PDF extraction unavailable")
        return ""
    pages: list[str] = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            text: Optional[str] = page.extract_text()
            if text:
                pages.append(text)
    return "\n\n".join(pages)


def _extract_docx(path: pathlib.Path) -> str:
    try:
        import docx  # type: ignore
    except ImportError:
        logger.warning("python-docx not installed — DOCX extraction unavailable")
        return ""
    doc = docx.Document(str(path))
    return "\n".join(para.text for para in doc.paragraphs if para.text.strip())


def _extract_xlsx(path: pathlib.Path) -> str:
    try:
        import openpyxl  # type: ignore
    except ImportError:
        logger.warning("openpyxl not installed — XLSX extraction unavailable")
        return ""
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets:
        lines.append(f"=== {sheet.title} ===")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            line = "\t".join(cells).strip()
            if line:
                lines.append(line)
    return "\n".join(lines)


def _extract_image_ocr(path: pathlib.Path) -> str:
    try:
        import pytesseract  # type: ignore
        from PIL import Image  # type: ignore
    except ImportError:
        logger.warning("pytesseract/Pillow not installed — OCR unavailable")
        return ""
    img = Image.open(str(path))
    return pytesseract.image_to_string(img, lang="dan+eng")
